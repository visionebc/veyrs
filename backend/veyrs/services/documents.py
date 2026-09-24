"""Document intelligence (spec section 18).

Pipeline: bytes -> text -> classification -> entity extraction -> correlation.

Two hard rules, both security rather than functionality:

* **Extraction never fetches anything.** An advisory HTML/XML document can
  reference external entities or remote images; resolving them from inside the
  parser turns "upload a PDF" into an SSRF primitive. XML is parsed with entity
  resolution disabled and no network access, and HTML is stripped textually.
* **Upload does not mean trust.** Extracted CVE ids are validated against the
  CVE table before anything is correlated; a document cannot invent a
  vulnerability, only reference one.

Heavy optional parsers (PDF/DOCX) are imported lazily so the core service runs
in a minimal container; a missing parser produces an explicit
`unsupported_format` status rather than a stack trace.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import logging
import re
import uuid
from typing import Any, Iterable

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Cve, Document, DocumentChunk, Product, Vendor
from . import correlation
from .versions import normalize_name

log = logging.getLogger("veyrs.documents")

CVE_RE = re.compile(r"\bCVE-\d{4}-\d{4,7}\b", re.IGNORECASE)
CWE_RE = re.compile(r"\bCWE-\d{1,4}\b", re.IGNORECASE)
CVSS_VECTOR_RE = re.compile(
    r"\bCVSS:(?:3\.[01]|4\.0)/[A-Z]{1,3}:[A-Z](?:/[A-Z]{1,3}:[A-Z])+", re.IGNORECASE
)
CVSS_SCORE_RE = re.compile(
    r"\bCVSS(?:\s*v?[234](?:\.\d)?)?\s*(?:base\s*)?score[:\s]+(\d{1,2}(?:\.\d)?)",
    re.IGNORECASE,
)
VERSION_RE = re.compile(r"\b\d+\.\d+(?:\.\d+){0,2}\b")
SEVERITY_RE = re.compile(r"\b(critical|high|medium|moderate|low|informational)\b", re.IGNORECASE)
# Advisories phrase the fix a dozen ways: "fixed in 7.2.5", "upgrade to FortiWeb
# version 7.2.5 or above", "please update to release 7.2.5". Allow a short run of
# non-numeric words (product name, "version", "release") between the keyword and
# the number, but not enough to reach an unrelated version elsewhere in the text.
FIXED_RE = re.compile(
    r"(?:fixed|patched|resolved|remediated|upgrade|update|migrate)\s+"
    r"(?:in|to)\s+(?:[A-Za-z][\w.-]*\s+){0,4}"
    r"([0-9]+\.[0-9]+(?:\.[0-9]+){0,2})",
    re.IGNORECASE,
)

MAX_TEXT_CHARS = 4_000_000  # ~4 MB of text; beyond this a "document" is a dump

DOC_CLASSES = ("vendor_advisory", "scan_report", "policy", "runbook", "news", "other")


class ExtractionError(RuntimeError):
    pass


def content_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------


def extract_text(data: bytes, content_type: str, filename: str = "") -> tuple[str, dict]:
    """Return (text, metadata). Raises ExtractionError for unsupported input."""
    kind = (content_type or "").split(";")[0].strip().lower()
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if kind in ("text/plain", "text/markdown") or suffix in ("txt", "md", "log"):
        return _decode(data), {}
    if kind == "application/json" or suffix == "json":
        return _from_json(data), {}
    if kind in ("text/csv", "application/csv") or suffix == "csv":
        return _from_csv(data), {}
    if kind in ("text/html", "application/xhtml+xml") or suffix in ("html", "htm"):
        return _from_html(_decode(data)), {}
    if kind in ("application/xml", "text/xml") or suffix in ("xml", "nessus"):
        return _from_xml(data), {}
    if kind == "application/pdf" or suffix == "pdf":
        return _from_pdf(data)
    if suffix == "docx" or kind == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        return _from_docx(data)
    raise ExtractionError(f"unsupported document type: {content_type or suffix or 'unknown'}")


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)[:MAX_TEXT_CHARS]
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")[:MAX_TEXT_CHARS]


def _from_json(data: bytes) -> str:
    try:
        parsed = json.loads(_decode(data))
    except json.JSONDecodeError as exc:
        raise ExtractionError(f"invalid JSON: {exc}") from exc
    return json.dumps(parsed, indent=2, ensure_ascii=False)[:MAX_TEXT_CHARS]


def _from_csv(data: bytes) -> str:
    text = _decode(data)
    reader = csv.reader(io.StringIO(text))
    lines = [" | ".join(row) for row in reader]
    return "\n".join(lines)[:MAX_TEXT_CHARS]


_TAG_RE = re.compile(r"<[^>]+>")
_SCRIPT_RE = re.compile(r"<(script|style)[^>]*>.*?</\1>", re.IGNORECASE | re.DOTALL)


def _from_html(text: str) -> str:
    """Strip tags textually. No parser, therefore no external entity fetches."""
    text = _SCRIPT_RE.sub(" ", text)
    text = _TAG_RE.sub(" ", text)
    text = (text.replace("&nbsp;", " ").replace("&amp;", "&")
                .replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"'))
    return re.sub(r"[ \t]{2,}", " ", text).strip()[:MAX_TEXT_CHARS]


def _from_xml(data: bytes) -> str:
    """Parse XML with entity declarations refused (XXE defence).

    Same dead guard as `importers/parsers._parse_xml` had, same fix: the
    handlers were attached to `XMLParser.parser`, gone since Python 3.9, under
    an `except AttributeError: pass`.
    """
    from xml.etree import ElementTree

    from . import xmlsafe

    try:
        xmlsafe.reject_entity_declarations(data)
    except xmlsafe.UnsafeXmlError as exc:
        raise ExtractionError(str(exc)) from exc
    try:
        root = ElementTree.fromstring(_decode(data))
    except ElementTree.ParseError as exc:
        raise ExtractionError(f"invalid XML: {exc}") from exc

    parts: list[str] = []
    for element in root.iter():
        if element.text and element.text.strip():
            parts.append(element.text.strip())
        for value in (element.attrib or {}).values():
            if value and len(value) < 500:
                parts.append(str(value))
    return "\n".join(parts)[:MAX_TEXT_CHARS]


def _from_pdf(data: bytes) -> tuple[str, dict]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise ExtractionError("PDF support requires the 'pypdf' package") from exc
    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "") for page in reader.pages]
    text = "\n\n".join(pages)[:MAX_TEXT_CHARS]
    meta = {"page_count": len(pages)}
    if len(text.strip()) < 40 * max(len(pages), 1):
        # Almost no extractable text: a scanned document. OCR is a separate,
        # explicitly-enabled step; flag it instead of silently returning noise.
        meta["needs_ocr"] = True
    return text, meta


def _from_docx(data: bytes) -> tuple[str, dict]:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise ExtractionError("DOCX support requires the 'python-docx' package") from exc
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)[:MAX_TEXT_CHARS], {}


# --------------------------------------------------------------------------
# Classification + entity extraction
# --------------------------------------------------------------------------


def classify(text: str, filename: str = "") -> str:
    lowered = f"{filename}\n{text[:8000]}".lower()
    if any(k in lowered for k in ("psirt", "security advisory", "advisory id", "fg-ir-")):
        return "vendor_advisory"
    if any(k in lowered for k in ("nessus", "qualys", "openvas", "greenbone", "scan report",
                                  "plugin id")):
        return "scan_report"
    if any(k in lowered for k in ("runbook", "step 1", "procedure")):
        return "runbook"
    if any(k in lowered for k in ("policy", "iso 27001", "shall be", "must be enforced")):
        return "policy"
    if any(k in lowered for k in ("posted on", "read more", "by staff writer")):
        return "news"
    return "other"


def extract_entities(text: str, session: Session | None = None) -> dict[str, Any]:
    """Pull CVEs, CWEs, CVSS, versions, severities and fixes out of free text."""
    cve_ids = sorted({m.upper() for m in CVE_RE.findall(text)})
    cwe_ids = sorted({m.upper() for m in CWE_RE.findall(text)})
    vectors = sorted({m.upper() for m in CVSS_VECTOR_RE.findall(text)})

    scores = [float(m) for m in CVSS_SCORE_RE.findall(text)]
    scores = [s for s in scores if 0.0 <= s <= 10.0]

    severities = [s.lower() for s in SEVERITY_RE.findall(text)]
    severity = None
    for candidate in ("critical", "high", "medium", "moderate", "low", "informational"):
        if candidate in severities:
            severity = "medium" if candidate == "moderate" else candidate
            break

    fixed_versions = sorted({m for m in FIXED_RE.findall(text)})

    vendors: list[str] = []
    product_ids: list[str] = []
    if session is not None:
        vendors, product_ids = _match_products(text, session)

    return {
        "cve_ids": cve_ids,
        "cwe_ids": cwe_ids,
        "cvss_vectors": vectors,
        "cvss_scores": scores,
        "max_cvss": max(scores) if scores else None,
        "severity": severity,
        "fixed_versions": fixed_versions,
        "versions": sorted(set(VERSION_RE.findall(text)))[:50],
        "vendors": vendors,
        "product_ids": product_ids,
    }


def _match_products(text: str, session: Session) -> tuple[list[str], list[str]]:
    """Match known vendors/products by name.

    Matching against the registry (rather than guessing product names out of the
    text) is what keeps this honest: an advisory can only be linked to a product
    VEYRS already knows, so a typo cannot invent inventory.
    """
    haystack = normalize_name(text[:20000])
    vendors: list[str] = []
    product_ids: list[str] = []
    for vendor in session.execute(select(Vendor)).scalars().all():
        if vendor.normalized_name and vendor.normalized_name in haystack:
            vendors.append(vendor.name)
    for product in session.execute(select(Product)).scalars().all():
        if product.normalized_name and product.normalized_name in haystack:
            product_ids.append(str(product.id))
    return sorted(set(vendors)), sorted(set(product_ids))


# --------------------------------------------------------------------------
# Chunking (for vector search)
# --------------------------------------------------------------------------


def chunk_text(text: str, size: int = 1200, overlap: int = 150) -> list[str]:
    """Split on paragraph boundaries, with overlap so context is not cut mid-fact."""
    if not text:
        return []
    paragraphs = [p.strip() for p in re.split(r"\n{2,}", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(current) + len(paragraph) + 2 <= size:
            current = f"{current}\n\n{paragraph}" if current else paragraph
            continue
        if current:
            chunks.append(current)
        if len(paragraph) <= size:
            current = paragraph
            continue
        # A single oversized paragraph is hard-split with overlap.
        start = 0
        while start < len(paragraph):
            chunks.append(paragraph[start:start + size])
            start += max(size - overlap, 1)
        current = ""
    if current:
        chunks.append(current)
    return chunks


# --------------------------------------------------------------------------
# Ingestion
# --------------------------------------------------------------------------


def ingest(
    session: Session,
    organization_id: uuid.UUID,
    *,
    filename: str,
    data: bytes,
    content_type: str = "text/plain",
    source_url: str | None = None,
    uploaded_by_id: uuid.UUID | None = None,
    data_classification: str = "internal",
    correlate: bool = True,
) -> tuple[Document, bool]:
    """Store, extract, classify and (optionally) correlate a document."""
    digest = content_hash(data)
    existing = session.execute(
        select(Document).where(
            Document.organization_id == organization_id, Document.content_hash == digest
        )
    ).scalars().first()
    if existing is not None:
        return existing, False

    document = Document(
        organization_id=organization_id, filename=filename[:400],
        content_type=content_type[:120], source_url=source_url,
        size_bytes=len(data), content_hash=digest, uploaded_by_id=uploaded_by_id,
        data_classification=data_classification, status="processing",
    )
    session.add(document)
    session.flush()

    try:
        text, meta = extract_text(data, content_type, filename)
    except ExtractionError as exc:
        document.status = "unsupported_format"
        document.error = str(exc)
        session.flush()
        return document, True
    except Exception as exc:  # noqa: BLE001 - a malformed upload must not 500
        log.exception("document extraction failed for %s", filename)
        document.status = "failed"
        document.error = f"{type(exc).__name__}: {exc}"[:2000]
        session.flush()
        return document, True

    document.text_content = text
    document.page_count = meta.get("page_count")
    document.ocr_used = False
    document.doc_class = classify(text, filename)

    entities = extract_entities(text, session)
    document.extracted = entities
    # Only CVEs VEYRS actually knows are recorded: a document cannot invent one.
    known = session.execute(
        select(Cve.id).where(Cve.id.in_(entities["cve_ids"] or ["-"]))
    ).scalars().all()
    document.cve_ids = sorted(known)
    if meta.get("needs_ocr"):
        document.status = "needs_ocr"
    else:
        document.status = "processed"

    for ordinal, chunk in enumerate(chunk_text(text)):
        session.add(DocumentChunk(
            organization_id=organization_id, document_id=document.id,
            ordinal=ordinal, text=chunk, tokens=max(len(chunk) // 4, 1),
        ))
    session.flush()

    if correlate and document.cve_ids:
        correlation.correlate_batch(session, organization_id, document.cve_ids)

    return document, True


def unknown_cves(document: Document) -> list[str]:
    """CVEs the document mentions that VEYRS has never seen.

    Surfaced in the UI as "intelligence gap" rather than dropped silently: it
    usually means the NVD feed is behind, which is itself worth knowing.
    """
    mentioned = set((document.extracted or {}).get("cve_ids") or [])
    return sorted(mentioned - set(document.cve_ids or []))


__all__ = [
    "ExtractionError", "extract_text", "extract_entities", "classify", "chunk_text",
    "ingest", "content_hash", "unknown_cves", "DOC_CLASSES",
]
